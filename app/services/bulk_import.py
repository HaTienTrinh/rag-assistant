import csv
import io
import os
import re
import zipfile
from concurrent.futures import ThreadPoolExecutor
from typing import Any

from fastapi import HTTPException
import openpyxl
from sqlalchemy import func, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.auth import hash_password
from app.config import settings
from app.models import User
from app.services.passwords import generate_temp_password

EMAIL_PATTERN = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
HEADER_KEYWORDS = {"email", "e-mail", "mail", "địa chỉ email"}

MAX_FILE_BYTES = 2 * 1024 * 1024  # 2 MB
MAX_UNCOMPRESSED_BYTES = 20 * 1024 * 1024  # 20 MB chống zip bomb


def generate_template_xlsx() -> bytes:
    """Tạo file Excel mẫu mau-tao-tai-khoan.xlsx trong bộ nhớ."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Danh sách tài khoản"

    ws["A1"] = "email"
    ws["A2"] = "nguyenvana@example.com"
    ws["A3"] = "tranthib@example.com"

    ws.column_dimensions["A"].width = 30

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _detect_csv_delimiter(sample_text: str) -> str:
    """Tự động phát hiện dấu phân cách phẩy (,) hoặc chấm phẩy (;) cho CSV."""
    first_line = sample_text.splitlines()[0] if sample_text else ""
    if ";" in first_line and first_line.count(";") >= first_line.count(","):
        return ";"
    return ","


def _extract_rows_from_xlsx(content: bytes) -> list[tuple[int, str]]:
    """Đọc dữ liệu từ file .xlsx và trích xuất danh sách (số_dòng, email_thô)."""
    # 1. Kiểm tra chữ ký ZIP (PK)
    if not content.startswith(b"PK"):
        raise HTTPException(status_code=400, detail="File không đúng định dạng Excel (.xlsx hợp lệ)")

    # 2. Kiểm tra tổng dung lượng giải nén chống zip bomb
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as zf:
            total_uncompressed = sum(info.file_size for info in zf.infolist())
            if total_uncompressed > MAX_UNCOMPRESSED_BYTES:
                raise HTTPException(
                    status_code=400,
                    detail=f"Dung lượng sau giải nén của file vượt quá giới hạn an toàn (tối đa {MAX_UNCOMPRESSED_BYTES // (1024 * 1024)} MB)",
                )
    except zipfile.BadZipFile:
        raise HTTPException(status_code=400, detail="File Excel bị hỏng hoặc không đúng định dạng ZIP")

    # 3. Mở workbook chỉ đọc bằng openpyxl
    try:
        wb = openpyxl.load_workbook(filename=io.BytesIO(content), read_only=True, data_only=True)
        wb_raw = openpyxl.load_workbook(filename=io.BytesIO(content), read_only=True, data_only=False)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Không thể đọc file Excel: {str(e)}")

    if not wb.worksheets or not wb_raw.worksheets:
        raise HTTPException(status_code=400, detail="File Excel không có sheet nào")

    ws = wb.worksheets[0]
    ws_raw = wb_raw.worksheets[0]

    raw_matrix: list[tuple[int, list[Any]]] = []
    for row_idx, (row, row_raw) in enumerate(
        zip(ws.iter_rows(values_only=True), ws_raw.iter_rows(values_only=True)), start=1
    ):
        if not row and not row_raw:
            continue
        # Kết hợp giá trị đã tính toán và giá trị công thức thô nếu ô chưa được tính
        len_max = max(len(row or ()), len(row_raw or ()))
        merged_row = []
        for i in range(len_max):
            v = row[i] if row and i < len(row) else None
            v_raw = row_raw[i] if row_raw and i < len(row_raw) else None
            merged_row.append(v if v is not None else v_raw)

        # Bỏ qua dòng trống hoàn toàn
        if all(c is None or str(c).strip() == "" for c in merged_row):
            continue
        raw_matrix.append((row_idx, merged_row))

    wb.close()
    wb_raw.close()

    if not raw_matrix:
        raise HTTPException(status_code=400, detail="File Excel không có dữ liệu")

    # 4. Tìm cột email ở dòng dữ liệu đầu tiên
    first_row_idx, first_row_cells = raw_matrix[0]
    email_col_idx = None
    has_header = False

    for col_idx, cell_val in enumerate(first_row_cells):
        if cell_val is None:
            continue
        val_norm = str(cell_val).strip().lower()
        # Bỏ khoảng trắng dư thừa
        val_norm = re.sub(r"\s+", " ", val_norm)
        if val_norm in HEADER_KEYWORDS:
            email_col_idx = col_idx
            has_header = True
            break

    if email_col_idx is None:
        # Không có tiêu đề khớp -> dùng cột A (index 0) và coi dòng đầu là dữ liệu
        email_col_idx = 0
        has_header = False

    data_rows: list[tuple[int, str]] = []
    start_offset = 1 if has_header else 0

    for row_idx, cells in raw_matrix[start_offset:]:
        val = cells[email_col_idx] if email_col_idx < len(cells) else ""
        data_rows.append((row_idx, str(val) if val is not None else ""))

    return data_rows


def _extract_rows_from_csv(content: bytes) -> list[tuple[int, str]]:
    """Đọc dữ liệu từ file .csv và trích xuất danh sách (số_dòng, email_thô)."""
    try:
        # utf-8-sig tự động xử lý và loại bỏ BOM nếu có
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise HTTPException(status_code=400, detail="File CSV không đúng định dạng mã hóa UTF-8")

    delimiter = _detect_csv_delimiter(text)
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)

    raw_matrix: list[tuple[int, list[str]]] = []
    for row_idx, row in enumerate(reader, start=1):
        if not row:
            continue
        if all(str(c).strip() == "" for c in row):
            continue
        raw_matrix.append((row_idx, row))

    if not raw_matrix:
        raise HTTPException(status_code=400, detail="File CSV không có dữ liệu")

    first_row_idx, first_row_cells = raw_matrix[0]
    email_col_idx = None
    has_header = False

    for col_idx, cell_val in enumerate(first_row_cells):
        val_norm = re.sub(r"\s+", " ", cell_val.strip().lower())
        if val_norm in HEADER_KEYWORDS:
            email_col_idx = col_idx
            has_header = True
            break

    if email_col_idx is None:
        email_col_idx = 0
        has_header = False

    data_rows: list[tuple[int, str]] = []
    start_offset = 1 if has_header else 0

    for row_idx, cells in raw_matrix[start_offset:]:
        val = cells[email_col_idx] if email_col_idx < len(cells) else ""
        data_rows.append((row_idx, val))

    return data_rows


def parse_and_classify_import_file(
    filename: str, content: bytes, db: Session
) -> dict[str, Any]:
    """Phân tích, chuẩn hóa và phân loại toàn bộ dòng trong file import."""
    # Kiểm tra kích thước file thô
    if len(content) > MAX_FILE_BYTES:
        raise HTTPException(
            status_code=400,
            detail=f"Kích thước file ({len(content) / (1024 * 1024):.2f} MB) vượt quá giới hạn cho phép (tối đa 2 MB)",
        )

    ext = os.path.splitext(filename)[1].lower()
    if ext == ".xlsx":
        raw_rows = _extract_rows_from_xlsx(content)
    elif ext == ".csv":
        raw_rows = _extract_rows_from_csv(content)
    else:
        raise HTTPException(
            status_code=400,
            detail="Định dạng file không được hỗ trợ. Vui lòng tải lên file Excel (.xlsx) hoặc CSV (.csv)",
        )

    total_rows = len(raw_rows)
    if total_rows == 0:
        raise HTTPException(status_code=400, detail="File không chứa dòng dữ liệu nào để nhập")

    # Giới hạn số dòng cấu hình bằng BULK_IMPORT_MAX_ROWS
    max_allowed_rows = settings.bulk_import_max_rows
    if total_rows > max_allowed_rows:
        raise HTTPException(
            status_code=413,
            detail=f"Số lượng dòng dữ liệu ({total_rows}) vượt quá giới hạn cho phép (tối đa {max_allowed_rows} dòng)",
        )

    # Thu thập tất cả email tiềm năng để truy vấn DB một lần
    candidate_emails = set()
    for _, raw_val in raw_rows:
        val_clean = str(raw_val).strip().lower()
        if val_clean and "@" in val_clean:
            candidate_emails.add(val_clean)

    # Truy vấn trạng thái tài khoản trong DB cho các email tiềm năng
    db_users_map: dict[str, Any] = {}
    if candidate_emails:
        query = select(User.email, User.deleted_at).where(
            func.lower(User.email).in_(list(candidate_emails))
        )
        for u_email, u_deleted_at in db.execute(query).all():
            db_users_map[u_email.lower()] = u_deleted_at

    classified_rows: list[dict[str, Any]] = []
    seen_in_file: set[str] = set()

    summary_counts = {
        "valid": 0,
        "created": 0,
        "invalid": 0,
        "duplicate_in_file": 0,
        "exists": 0,
        "in_trash": 0,
    }

    for row_idx, raw_val in raw_rows:
        val_str = str(raw_val).strip()
        norm_email = val_str.lower()

        # 1. Kiểm tra tính hợp lệ cơ bản
        if not val_str:
            status = "invalid"
            msg = "Ô email bị trống"
        elif "," in val_str or ";" in val_str:
            status = "invalid"
            msg = "Ô chứa nhiều email hoặc có dấu phân cách không hợp lệ"
        elif any(c.isspace() for c in val_str):
            status = "invalid"
            msg = "Email không được chứa khoảng trắng"
        elif len(norm_email) > 255:
            status = "invalid"
            msg = "Email quá dài (tối đa 255 ký tự)"
        elif not EMAIL_PATTERN.match(norm_email):
            status = "invalid"
            msg = "Định dạng email không hợp lệ"
        elif norm_email in seen_in_file:
            # 2. Trùng trong file (từ lần xuất hiện thứ 2 trở đi)
            status = "duplicate_in_file"
            msg = "Email bị trùng lặp trong file"
        elif norm_email in db_users_map:
            deleted_at = db_users_map[norm_email]
            if deleted_at is not None:
                # 3. Thuộc tài khoản trong thùng rác
                status = "in_trash"
                msg = "Email thuộc một tài khoản trong thùng rác, hãy khôi phục hoặc xóa vĩnh viễn"
            else:
                # 4. Đã tồn tại trong hệ thống
                status = "exists"
                msg = "Email đã tồn tại trong hệ thống"
        else:
            # 5. Hợp lệ
            status = "valid"
            msg = "Hợp lệ"

        if val_str and norm_email and EMAIL_PATTERN.match(norm_email):
            seen_in_file.add(norm_email)

        summary_counts[status] += 1
        classified_rows.append({
            "row": row_idx,
            "email": norm_email if status != "invalid" else val_str,
            "status": status,
            "message": msg,
        })

    # Đảm bảo sắp xếp theo số dòng Excel
    classified_rows.sort(key=lambda x: x["row"])

    return {
        "total_rows": total_rows,
        "summary": summary_counts,
        "rows": classified_rows,
    }


def execute_bulk_import_creation(
    classified_data: dict[str, Any],
    daily_limit: int | None,
    require_password_change: bool,
    db: Session,
) -> dict[str, Any]:
    """Tạo tài khoản thật cho các dòng 'valid' trong một transaction."""
    rows = classified_data["rows"]
    valid_rows = [r for r in rows if r["status"] == "valid"]

    if not valid_rows:
        return {
            "dry_run": False,
            "total_rows": classified_data["total_rows"],
            "summary": classified_data["summary"],
            "rows": rows,
        }

    # 1. Sinh mật khẩu tạm ngẫu nhiên cho từng dòng valid
    passwords = [generate_temp_password(12) for _ in valid_rows]

    # 2. Băm mật khẩu song song bằng ThreadPoolExecutor
    max_workers = min(8, os.cpu_count() or 4)
    with ThreadPoolExecutor(max_workers=max_workers) as executor:
        hashed_passwords = list(executor.map(hash_password, passwords))

    # 3. Tạo các đối tượng User và commit trong một transaction
    new_users = []
    for r, pwd_hash in zip(valid_rows, hashed_passwords):
        user = User(
            email=r["email"],
            password_hash=pwd_hash,
            role="user",
            is_active=True,
            daily_limit=daily_limit,
            must_change_password=require_password_change,
        )
        new_users.append(user)

    try:
        db.add_all(new_users)
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(
            status_code=409,
            detail="Có xung đột dữ liệu, vui lòng thử lại",
        )

    # 4. Gắn kết quả mật khẩu vào các dòng đã tạo
    for r, pwd in zip(valid_rows, passwords):
        r["status"] = "created"
        r["message"] = "Đã tạo tài khoản thành công"
        r["password"] = pwd

    summary = dict(classified_data["summary"])
    summary["created"] = len(valid_rows)
    summary["valid"] = 0

    return {
        "dry_run": False,
        "total_rows": classified_data["total_rows"],
        "summary": summary,
        "rows": rows,
    }
