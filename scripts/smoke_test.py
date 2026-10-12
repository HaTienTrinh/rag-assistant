"""Script kiểm thử toàn diện các kịch bản của hệ thống (Smoke Test).

Sử dụng:
    python scripts/smoke_test.py
    python scripts/smoke_test.py --base-url http://127.0.0.1:8000 --admin-email admin@example.com --admin-password Admin12345
"""

import argparse
import os
import sys
import io
import time
import uuid
from typing import Any

import httpx
import openpyxl

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8")

# Cho phép import config từ ứng dụng nếu có
try:
    from app.config import settings
    DEFAULT_ADMIN_EMAIL = getattr(settings, "admin_email", "") or "admin@example.com"
    DEFAULT_ADMIN_PASSWORD = getattr(settings, "admin_password", "") or "Admin12345"
except Exception:
    DEFAULT_ADMIN_EMAIL = "admin@example.com"
    DEFAULT_ADMIN_PASSWORD = "Admin12345"



class TestResult:
    def __init__(self):
        self.results: list[dict[str, Any]] = []

    def record(self, scenario: str, test_case: str, passed: bool, detail: str = ""):
        self.results.append({
            "scenario": scenario,
            "test_case": test_case,
            "passed": passed,
            "detail": detail
        })
        status_str = "[ĐẠT]" if passed else "[KHÔNG ĐẠT]"
        print(f"  {status_str} {test_case} {f'({detail})' if detail else ''}")

    def print_summary(self):
        print("\n" + "=" * 80)
        print("BẢNG TỔNG KẾT KẾT QUẢ KIỂM THỬ (SMOKE TEST)")
        print("=" * 80)
        print(f"{'Kịch bản':<35} | {'Trường hợp kiểm thử':<30} | {'Kết quả':<10}")
        print("-" * 80)
        passed_count = 0
        total_count = len(self.results)
        for r in self.results:
            status = "ĐẠT" if r["passed"] else "KHÔNG ĐẠT"
            if r["passed"]:
                passed_count += 1
            print(f"{r['scenario'][:34]:<35} | {r['test_case'][:29]:<30} | {status:<10}")
        print("-" * 80)
        print(f"Tổng kết: {passed_count}/{total_count} bài kiểm thử thành công.")
        print("=" * 80 + "\n")
        return passed_count == total_count


def auth_login(client: httpx.Client, email: str, password: str) -> httpx.Response:
    """Đăng nhập và tự động gắn access_token vào Authorization header."""
    resp = client.post("/auth/login", data={"username": email, "password": password})
    if resp.status_code == 200:
        token = resp.json().get("access_token")
        if token:
            client.headers["Authorization"] = f"Bearer {token}"
    return resp


def auth_refresh(client: httpx.Client) -> httpx.Response:
    """Refresh token và tự động cập nhật Authorization header."""
    resp = client.post("/auth/refresh")
    if resp.status_code == 200:
        token = resp.json().get("access_token")
        if token:
            client.headers["Authorization"] = f"Bearer {token}"
    return resp


def create_user_via_admin(
    client: httpx.Client, base_url: str, email: str, password: str, role: str = "user", require_change: bool = False
) -> int:
    """Tạo người dùng thông qua API quản trị."""
    resp = client.post(
        f"{base_url}/admin/users",
        json={"email": email, "password": password, "role": role, "daily_limit": None, "require_change": require_change}
    )
    if resp.status_code != 201:
        raise RuntimeError(f"Không thể tạo user thử nghiệm: {resp.status_code} {resp.text}")
    return resp.json()["id"]


def purge_user_via_admin(client: httpx.Client, base_url: str, user_id: int):
    """Xóa hoàn toàn người dùng (chuyển thùng rác rồi xóa vĩnh viễn)."""
    # 1. Chuyển vào thùng rác nếu chưa
    client.delete(f"{base_url}/admin/users/{user_id}")
    # 2. Xóa vĩnh viễn
    client.delete(f"{base_url}/admin/users/{user_id}/permanent")


def run_smoke_tests(base_url: str, admin_email: str, admin_password: str) -> bool:
    tr = TestResult()
    print(f"\nBắt đầu chạy Smoke Test trên {base_url}...")
    print(f"Tài khoản Admin: {admin_email}\n")

    admin_client = httpx.Client(base_url=base_url, timeout=15.0)
    users_to_cleanup: list[int] = []

    try:
        # ==========================================
        # 0. ĐĂNG NHẬP TÀI KHOẢN ADMIN
        # ==========================================
        print(">> Kiểm tra đăng nhập Admin...")
        login_resp = auth_login(admin_client, admin_email, admin_password)
        if login_resp.status_code != 200:
            print(f"LỖI: Không thể đăng nhập tài khoản admin ({login_resp.status_code}): {login_resp.text}")
            return False
        tr.record("Xác thực Admin", "Đăng nhập admin thành công", True)

        # ==========================================
        # KỊCH BẢN 1: ĐĂNG NHẬP, REFRESH XOAY VÒNG, REPLAY ATTACK, ĐĂNG XUẤT
        # ==========================================
        print("\n>> Kịch bản 1: Đăng nhập, refresh xoay vòng, từ chối token cũ, đăng xuất...")
        test_email = f"smoke_user_{uuid.uuid4().hex[:8]}@example.com"
        test_password = "UserPass12345@"
        user_id = create_user_via_admin(admin_client, base_url, test_email, test_password)
        users_to_cleanup.append(user_id)

        user_client = httpx.Client(base_url=base_url, timeout=15.0)
        # 1.1 Đăng nhập
        u_login = auth_login(user_client, test_email, test_password)
        tr.record("Kịch bản 1 (Auth)", "User đăng nhập thành công", u_login.status_code == 200)

        old_refresh_token = user_client.cookies.get("refresh_token")

        # 1.2 Refresh xoay vòng lần 1
        u_refresh1 = auth_refresh(user_client)
        new_refresh_token = user_client.cookies.get("refresh_token")
        rotated = (u_refresh1.status_code == 200) and (new_refresh_token != old_refresh_token)
        tr.record("Kịch bản 1 (Auth)", "Refresh xoay vòng cấp token mới", rotated)

        # 1.3 Dùng lại refresh token cũ sau khi bị thay thế và sau thời gian ân hạn
        attacker_client = httpx.Client(base_url=base_url, timeout=15.0)
        attacker_client.cookies.set("refresh_token", old_refresh_token)
        # Thử refresh lần 2 từ user hợp lệ để token cũ hoàn toàn bị coi là cũ/thay thế
        auth_refresh(user_client)
        
        # Chờ 11s để vượt qua REFRESH_REUSE_GRACE_SECONDS (10s) và kích hoạt cơ chế phát hiện tái sử dụng
        print("  (Chờ 11s để hết thời gian ân hạn và kích hoạt cơ chế phát hiện tái sử dụng token...)")
        time.sleep(11.0)
        attack_resp = attacker_client.post("/auth/refresh")
        old_token_rejected = (attack_resp.status_code in (401, 403))
        tr.record("Kịch bản 1 (Auth)", "Dùng lại refresh token cũ bị từ chối/thu hồi", old_token_rejected)

        # 1.4 Đăng xuất làm refresh thất bại
        auth_login(user_client, test_email, test_password)
        u_logout = user_client.post("/auth/logout")
        tr.record("Kịch bản 1 (Auth)", "Đăng xuất thành công", u_logout.status_code == 200)

        # Thử refresh sau khi đăng xuất
        u_refresh_after_logout = user_client.post("/auth/refresh")
        tr.record("Kịch bản 1 (Auth)", "Refresh thất bại sau khi đã đăng xuất", u_refresh_after_logout.status_code in (401, 403))

        # ==========================================
        # KỊCH BẢN 2: ĐỔI MẬT KHẨU, THU HỒI PHIÊN KHÁC
        # ==========================================
        print("\n>> Kịch bản 2: Đổi mật khẩu, thu hồi phiên khác, kiểm tra mật khẩu cũ...")
        new_password = "NewSecretPass6789@"
        
        # Mở phiên A (Client A)
        session_a = httpx.Client(base_url=base_url, timeout=15.0)
        auth_login(session_a, test_email, test_password)

        # Mở phiên B (Client B)
        session_b = httpx.Client(base_url=base_url, timeout=15.0)
        auth_login(session_b, test_email, test_password)

        # Phiên A đổi mật khẩu
        chg_resp = session_a.post(
            "/auth/change-password",
            json={"current_password": test_password, "new_password": new_password}
        )
        tr.record("Kịch bản 2 (Password)", "Đổi mật khẩu thành công", chg_resp.status_code == 200)

        # Phiên A nhận access_token mới trả về từ change-password và cập nhật header
        if chg_resp.status_code == 200 and chg_resp.json().get("access_token"):
            session_a.headers["Authorization"] = f"Bearer {chg_resp.json()['access_token']}"

        # Phiên A vẫn truy cập được
        me_a = session_a.get("/auth/me")
        tr.record("Kịch bản 2 (Password)", "Phiên hiện tại (A) giữ nguyên hoạt động", me_a.status_code == 200)

        # Phiên B bị thu hồi (refresh thất bại)
        refresh_b = session_b.post("/auth/refresh")
        tr.record("Kịch bản 2 (Password)", "Phiên khác (B) bị thu hồi khi refresh", refresh_b.status_code in (401, 403))

        # Đăng nhập lại với mật khẩu cũ phải thất bại
        old_login = httpx.post(f"{base_url}/auth/login", data={"username": test_email, "password": test_password})
        tr.record("Kịch bản 2 (Password)", "Mật khẩu cũ không đăng nhập được", old_login.status_code == 401)

        # Đăng nhập với mật khẩu mới phải thành công
        new_login = httpx.post(f"{base_url}/auth/login", data={"username": test_email, "password": new_password})
        tr.record("Kịch bản 2 (Password)", "Mật khẩu mới đăng nhập thành công", new_login.status_code == 200)


        # ==========================================
        # KỊCH BẢN 3: QUẢN TRỊ USER, KHÓA/MỞ, THÙNG RÁC, CHẶN ADMIN CUỐI
        # ==========================================
        print("\n>> Kịch bản 3: Khóa, mở khóa, thùng rác, khôi phục, xóa vĩnh viễn, trùng email, chặn admin cuối...")
        # 3.1 Khóa tài khoản
        lock_resp = admin_client.patch(f"/admin/users/{user_id}", json={"is_active": False})
        tr.record("Kịch bản 3 (User)", "Khóa tài khoản thành công", lock_resp.status_code == 200 and lock_resp.json()["is_active"] is False)

        # Tài khoản bị khóa không đăng nhập được
        locked_login = httpx.post(f"{base_url}/auth/login", data={"username": test_email, "password": new_password})
        tr.record("Kịch bản 3 (User)", "Tài khoản bị khóa bị từ chối đăng nhập", locked_login.status_code in (401, 403))

        # 3.2 Mở khóa tài khoản
        unlock_resp = admin_client.patch(f"/admin/users/{user_id}", json={"is_active": True})
        tr.record("Kịch bản 3 (User)", "Mở khóa tài khoản thành công", unlock_resp.status_code == 200 and unlock_resp.json()["is_active"] is True)

        # 3.3 Chuyển vào thùng rác
        trash_resp = admin_client.delete(f"/admin/users/{user_id}")
        tr.record("Kịch bản 3 (User)", "Chuyển tài khoản vào thùng rác", trash_resp.status_code == 200)

        # 3.4 Trùng email trong thùng rác -> 409
        dup_resp = admin_client.post(
            "/admin/users",
            json={"email": test_email, "password": "SamplePassword123@", "role": "user"}
        )
        tr.record("Kịch bản 3 (User)", "Trùng email tài khoản trong thùng rác trả về 409", dup_resp.status_code == 409)

        # 3.5 Khôi phục tài khoản từ thùng rác
        restore_resp = admin_client.post(f"/admin/users/{user_id}/restore")
        tr.record("Kịch bản 3 (User)", "Khôi phục tài khoản từ thùng rác", restore_resp.status_code == 200)

        # 3.6 Chặn tự khóa hoặc xóa admin cuối cùng
        me_admin = admin_client.get("/auth/me").json()
        admin_id = me_admin["id"]

        # Thử tự khóa chính admin
        self_lock = admin_client.patch(f"/admin/users/{admin_id}", json={"is_active": False})
        tr.record("Kịch bản 3 (User)", "Chặn tự khóa tài khoản của chính admin", self_lock.status_code in (400, 409))

        # Thử chuyển chính admin vào thùng rác
        self_trash = admin_client.delete(f"/admin/users/{admin_id}")
        tr.record("Kịch bản 3 (User)", "Chặn chuyển chính admin vào thùng rác", self_trash.status_code in (400, 409))

        # 3.7 Xóa vĩnh viễn user sau khi chuyển vào thùng rác
        admin_client.delete(f"/admin/users/{user_id}")  # vào thùng rác
        perm_del = admin_client.delete(f"/admin/users/{user_id}/permanent")
        tr.record("Kịch bản 3 (User)", "Xóa vĩnh viễn tài khoản trong thùng rác", perm_del.status_code == 200)
        # Đã xóa vĩnh viễn xong, bỏ khỏi danh sách cleanup
        users_to_cleanup.remove(user_id)

        # ==========================================
        # KỊCH BẢN 4: PHÂN QUYỀN VÀ BẢO VỆ DỮ LIỆU
        # ==========================================
        print("\n>> Kịch bản 4: Phân quyền admin/user và bảo vệ quyền sở hữu hội thoại...")
        user_a_email = f"user_a_{uuid.uuid4().hex[:8]}@example.com"
        user_b_email = f"user_b_{uuid.uuid4().hex[:8]}@example.com"
        pass_ab = "StandardUser123@"

        user_a_id = create_user_via_admin(admin_client, base_url, user_a_email, pass_ab)
        users_to_cleanup.append(user_a_id)
        user_b_id = create_user_via_admin(admin_client, base_url, user_b_email, pass_ab)
        users_to_cleanup.append(user_b_id)

        client_a = httpx.Client(base_url=base_url, timeout=15.0)
        auth_login(client_a, user_a_email, pass_ab)

        client_b = httpx.Client(base_url=base_url, timeout=15.0)
        auth_login(client_b, user_b_email, pass_ab)

        # 4.1 User thường gọi /admin/* -> 403
        a_admin_users = client_a.get("/admin/users")
        tr.record("Kịch bản 4 (Authz)", "User thường gọi /admin/users trả về 403", a_admin_users.status_code == 403)

        # 4.2 User thường gọi /documents -> 403
        a_docs = client_a.get("/documents")
        tr.record("Kịch bản 4 (Authz)", "User thường gọi /documents trả về 403", a_docs.status_code == 403)

        # 4.3 User thường gọi /rag/search -> 403
        a_rag = client_a.post("/rag/search", json={"query": "test"})
        tr.record("Kịch bản 4 (Authz)", "User thường gọi /rag/search trả về 403", a_rag.status_code == 403)


        # 4.4 User B tạo hội thoại, User A thử truy cập -> 404
        conv_b_resp = client_b.post("/conversations", json={"title": "Hội thoại của User B"})
        conv_b_id = conv_b_resp.json()["id"]

        a_read_b = client_a.get(f"/conversations/{conv_b_id}/messages")
        tr.record("Kịch bản 4 (Authz)", "User A đọc hội thoại của User B trả về 404", a_read_b.status_code == 404)

        a_ask_b = client_a.post(f"/conversations/{conv_b_id}/ask", json={"question": "Câu hỏi trái phép"})
        tr.record("Kịch bản 4 (Authz)", "User A gửi câu hỏi vào hội thoại của User B trả về 404", a_ask_b.status_code == 404)

        a_del_b = client_a.delete(f"/conversations/{conv_b_id}")
        tr.record("Kịch bản 4 (Authz)", "User A xóa hội thoại của User B trả về 404", a_del_b.status_code == 404)

        # ==========================================
        # KỊCH BẢN 5: THỐNG KÊ /admin/stats
        # ==========================================
        print("\n>> Kịch bản 5: Kiểm tra /admin/stats theo các chu kỳ và tham số sai...")
        for period in ["recent", "week", "month", "year"]:
            stats_resp = admin_client.get(f"/admin/stats?period={period}")
            valid = (stats_resp.status_code == 200) and ("buckets" in stats_resp.json()) and ("totals" in stats_resp.json())
            tr.record("Kịch bản 5 (Stats)", f"Thống kê /admin/stats period={period}", valid)

        # Tham số sai
        stats_bad_period = admin_client.get("/admin/stats?period=invalid_period")
        tr.record("Kịch bản 5 (Stats)", "Gọi period không hợp lệ trả về 400 hoặc 422", stats_bad_period.status_code in (400, 422))

        stats_bad_days = admin_client.get("/admin/stats?period=recent&days=9999")
        tr.record("Kịch bản 5 (Stats)", "Gọi days không hợp lệ (>365) trả về 400 hoặc 422", stats_bad_days.status_code in (400, 422))

        # ==========================================
        # KỊCH BẢN 6: BẮT BUỘC ĐỔI MẬT KHẨU (PHASE IMPORT-A)
        # ==========================================
        print("\n>> Kịch bản 6: Bắt buộc đổi mật khẩu ở lần đăng nhập đầu tiên (PHASE IMPORT-A)...")
        # 6.1 Tạo user mặc định (không truyền require_change) -> mặc định phải là require_change=True
        req_email = f"must_change_{uuid.uuid4().hex[:8]}@example.com"
        req_pass_init = "InitialPass123@"
        resp_def = admin_client.post(
            f"{base_url}/admin/users",
            json={"email": req_email, "password": req_pass_init, "role": "user", "daily_limit": None}
        )
        tr.record("Kịch bản 6 (MustChange)", "Admin tạo user mặc định thành công", resp_def.status_code == 201)
        req_user_id = resp_def.json()["id"]
        users_to_cleanup.append(req_user_id)
        tr.record("Kịch bản 6 (MustChange)", "Trường must_change_password mặc định là True", resp_def.json().get("must_change_password") is True)

        # Kiểm tra GET /admin/users trả về must_change_password cho từng người
        admin_users_list = admin_client.get(f"{base_url}/admin/users").json()
        target_in_list = next((u for u in admin_users_list if u["id"] == req_user_id), None)
        tr.record("Kịch bản 6 (MustChange)", "GET /admin/users chứa cờ must_change_password=True", target_in_list is not None and target_in_list.get("must_change_password") is True)

        # 6.2 User đăng nhập bằng mật khẩu tạm
        req_client = httpx.Client(base_url=base_url, timeout=15.0)
        u_login = auth_login(req_client, req_email, req_pass_init)
        tr.record("Kịch bản 6 (MustChange)", "User đăng nhập bằng mật khẩu tạm thành công", u_login.status_code == 200)

        # 6.3 GET /auth/me trả về must_change_password=True
        me_resp = req_client.get("/auth/me")
        tr.record("Kịch bản 6 (MustChange)", "GET /auth/me trả về must_change_password=True", me_resp.status_code == 200 and me_resp.json().get("must_change_password") is True)

        # 6.4 Chặn gọi API nghiệp vụ (GET /conversations) với 403 và header X-Error-Code
        blocked_conv = req_client.get("/conversations")
        tr.record(
            "Kịch bản 6 (MustChange)",
            "Chặn API /conversations trả về 403 kèm header X-Error-Code",
            blocked_conv.status_code == 403
            and blocked_conv.headers.get("X-Error-Code") == "PASSWORD_CHANGE_REQUIRED"
            and "Bạn cần đổi mật khẩu trước khi tiếp tục" in blocked_conv.text
        )

        # 6.5 POST /auth/logout vẫn hoạt động bình thường khi có cờ
        logout_resp = req_client.post("/auth/logout")
        tr.record("Kịch bản 6 (MustChange)", "POST /auth/logout hoạt động khi cờ bật", logout_resp.status_code == 200)

        # Đăng nhập lại để tiếp tục đổi mật khẩu
        auth_login(req_client, req_email, req_pass_init)

        # 6.6 POST /auth/change-password thành công thì xóa cờ
        req_pass_new = "NewSecretPass888@"
        change_resp = req_client.post(
            "/auth/change-password",
            json={"current_password": req_pass_init, "new_password": req_pass_new}
        )
        tr.record("Kịch bản 6 (MustChange)", "Đổi mật khẩu thành công khi cờ bật", change_resp.status_code == 200)

        if change_resp.status_code == 200 and change_resp.json().get("access_token"):
            req_client.headers["Authorization"] = f"Bearer {change_resp.json()['access_token']}"

        # 6.7 Sau khi đổi xong: GET /auth/me trả về must_change_password=False
        me_after_change = req_client.get("/auth/me")
        tr.record("Kịch bản 6 (MustChange)", "Sau khi đổi: cờ must_change_password=False", me_after_change.status_code == 200 and me_after_change.json().get("must_change_password") is False)

        # 6.8 Sau khi đổi xong: API /conversations hoạt động bình thường
        unblocked_conv = req_client.get("/conversations")
        tr.record("Kịch bản 6 (MustChange)", "Sau khi đổi: API /conversations được phép truy cập", unblocked_conv.status_code == 200)

        # 6.9 Admin đặt lại mật khẩu với require_change=True
        req_pass_reset = "AdminResetPass999@"
        reset_resp = admin_client.post(
            f"{base_url}/admin/users/{req_user_id}/reset-password",
            json={"password": req_pass_reset, "require_change": True}
        )
        tr.record("Kịch bản 6 (MustChange)", "Admin đặt lại mật khẩu với require_change=True", reset_resp.status_code == 200)

        # Phiên cũ bị thu hồi
        old_session_blocked = req_client.get("/conversations")
        tr.record("Kịch bản 6 (MustChange)", "Đặt lại mật khẩu thu hồi phiên cũ", old_session_blocked.status_code in (401, 403))

        # Đăng nhập lại bằng mật khẩu mới đặt lại
        auth_login(req_client, req_email, req_pass_reset)
        me_after_reset = req_client.get("/auth/me")
        tr.record("Kịch bản 6 (MustChange)", "Đăng nhập sau khi đặt lại: cờ bật trở lại", me_after_reset.status_code == 200 and me_after_reset.json().get("must_change_password") is True)

        # Đổi lại mật khẩu để tài khoản trở về bình thường
        req_pass_final = "FinalNewPass777@"
        req_client.post(
            "/auth/change-password",
            json={"current_password": req_pass_reset, "new_password": req_pass_final}
        )

        # 6.10 Admin đặt lại mật khẩu với require_change=False
        req_pass_reset2 = "AdminResetPassNoChange1@"
        reset_resp2 = admin_client.post(
            f"{base_url}/admin/users/{req_user_id}/reset-password",
            json={"password": req_pass_reset2, "require_change": False}
        )
        tr.record("Kịch bản 6 (MustChange)", "Admin đặt lại mật khẩu với require_change=False", reset_resp2.status_code == 200)

        auth_login(req_client, req_email, req_pass_reset2)
        me_no_change = req_client.get("/auth/me")
        tr.record("Kịch bản 6 (MustChange)", "require_change=False thì không bật cờ", me_no_change.status_code == 200 and me_no_change.json().get("must_change_password") is False)

        conv_no_change = req_client.get("/conversations")
        tr.record("Kịch bản 6 (MustChange)", "Dùng bình thường ngay khi tắt checkbox", conv_no_change.status_code == 200)

        # 6.11 Admin bị bật cờ: cũng bị chặn truy cập /admin/* cho tới khi đổi
        admin_test_email = f"admin_must_change_{uuid.uuid4().hex[:8]}@example.com"
        admin_test_pass = "AdminPassInit123@"
        admin_test_id = create_user_via_admin(admin_client, base_url, admin_test_email, admin_test_pass, role="admin", require_change=True)
        users_to_cleanup.append(admin_test_id)

        admin_test_client = httpx.Client(base_url=base_url, timeout=15.0)
        auth_login(admin_test_client, admin_test_email, admin_test_pass)
        admin_blocked = admin_test_client.get("/admin/users")
        tr.record(
            "Kịch bản 6 (MustChange)",
            "Admin có cờ must_change_password bị chặn /admin/users trả về 403",
            admin_blocked.status_code == 403 and admin_blocked.headers.get("X-Error-Code") == "PASSWORD_CHANGE_REQUIRED"
        )

        # ==========================================
        # 7. KỊCH BẢN 7: TẠO TÀI KHOẢN HÀNG LOẠT (PHASE IMPORT-B)
        # ==========================================
        print("\n>> Kịch bản 7: Nhập tài khoản hàng loạt (PHASE IMPORT-B)...")

        # 7.1 Tải file mẫu
        tpl_resp = admin_client.get(f"{base_url}/admin/users-import/template")
        tr.record(
            "Kịch bản 7 (BulkImport)",
            "Tải file Excel mẫu mau-tao-tai-khoan.xlsx",
            tpl_resp.status_code == 200 and "application/vnd.openxmlformats" in tpl_resp.headers.get("content-type", "") and "mau-tao-tai-khoan.xlsx" in tpl_resp.headers.get("content-disposition", "")
        )

        try:
            tpl_wb = openpyxl.load_workbook(io.BytesIO(tpl_resp.content))
            sheet = tpl_wb.active
            is_valid_tpl = sheet["A1"].value == "email" and sheet["A2"].value and "example.com" in str(sheet["A2"].value)
        except Exception:
            is_valid_tpl = False
        tr.record("Kịch bản 7 (BulkImport)", "File mẫu đọc được bằng openpyxl và có cột email", is_valid_tpl)

        # 7.2 Phân quyền: Khách và User thường bị chặn
        unauth_client = httpx.Client(base_url=base_url, timeout=15.0)
        unauth_tpl = unauth_client.get(f"{base_url}/admin/users-import/template")
        unauth_imp = unauth_client.post(f"{base_url}/admin/users-import")
        tr.record("Kịch bản 7 (BulkImport)", "Chưa đăng nhập gọi endpoint nhận 401", unauth_tpl.status_code == 401 and unauth_imp.status_code == 401)

        user_tpl = req_client.get(f"{base_url}/admin/users-import/template")
        user_imp = req_client.post(f"{base_url}/admin/users-import")
        tr.record("Kịch bản 7 (BulkImport)", "User thường gọi endpoint nhận 403", user_tpl.status_code == 403 and user_imp.status_code == 403)

        # 7.3 Xác thực file lỗi & zip bomb / kích thước
        fake_xlsx_resp = admin_client.post(
            f"{base_url}/admin/users-import",
            data={"dry_run": "true"},
            files={"file": ("fake.xlsx", b"this is not a zip pk file", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        )
        tr.record("Kịch bản 7 (BulkImport)", "File .txt đổi tên thành .xlsx bị từ chối 400", fake_xlsx_resp.status_code == 400)

        big_file_resp = admin_client.post(
            f"{base_url}/admin/users-import",
            data={"dry_run": "true"},
            files={"file": ("big.xlsx", b"0" * (2 * 1024 * 1024 + 100), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        )
        tr.record("Kịch bản 7 (BulkImport)", "File vượt quá 2 MB bị từ chối 400", big_file_resp.status_code == 400)

        # 7.4 Phân loại đầy đủ với dry_run
        trash_email = f"trash_bulk_{uuid.uuid4().hex[:8]}@example.com"
        trash_uid = create_user_via_admin(admin_client, base_url, trash_email, "TrashPass1234@")
        admin_client.delete(f"{base_url}/admin/users/{trash_uid}")
        users_to_cleanup.append(trash_uid)

        users_before_dry = len(admin_client.get(f"{base_url}/admin/users").json())

        wb_dry = openpyxl.Workbook()
        ws_dry = wb_dry.active
        ws_dry.append(["Địa chỉ Email", "Ghi chú"])
        val_email = f"bulk_val_{uuid.uuid4().hex[:8]}@example.com"
        ws_dry.append([val_email, "Hợp lệ"])
        ws_dry.append([val_email.upper(), "Trùng lặp trong file"])
        ws_dry.append(["invalid_email_format", "Sai định dạng"])
        ws_dry.append([admin_email, "Đã tồn tại"])
        ws_dry.append([trash_email, "Trong thùng rác"])
        ws_dry.append(["multi1@example.com, multi2@example.com", "Nhiều email trong 1 ô"])
        ws_dry.append(["=1+1", "Công thức"])

        buf_dry = io.BytesIO()
        wb_dry.save(buf_dry)

        dry_resp = admin_client.post(
            f"{base_url}/admin/users-import",
            data={"dry_run": "true"},
            files={"file": ("test_classify.xlsx", buf_dry.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        )
        tr.record("Kịch bản 7 (BulkImport)", "Gọi dry_run thành công 200", dry_resp.status_code == 200)

        dry_json = dry_resp.json()
        summary = dry_json.get("summary", {})
        is_classified_correctly = (
            summary.get("valid") == 1
            and summary.get("duplicate_in_file") == 1
            and summary.get("invalid") == 3
            and summary.get("exists") == 1
            and summary.get("in_trash") == 1
            and summary.get("created") == 0
        )
        tr.record("Kịch bản 7 (BulkImport)", "Phân loại chính xác 5 nhóm trạng thái", is_classified_correctly, f"summary={summary}")

        users_after_dry = len(admin_client.get(f"{base_url}/admin/users").json())
        tr.record("Kịch bản 7 (BulkImport)", "dry_run không tạo tài khoản vào DB", users_before_dry == users_after_dry)

        # 7.5 Tạo thật (dry_run = False)
        wb_real = openpyxl.Workbook()
        ws_real = wb_real.active
        ws_real.append(["email"])
        real_user1_email = f"bulk_u1_{uuid.uuid4().hex[:8]}@example.com"
        real_user2_email = f"bulk_u2_{uuid.uuid4().hex[:8]}@example.com"
        ws_real.append([real_user1_email])
        ws_real.append([real_user2_email])

        buf_real = io.BytesIO()
        wb_real.save(buf_real)

        real_resp = admin_client.post(
            f"{base_url}/admin/users-import",
            data={"dry_run": "false", "daily_limit": "30", "require_password_change": "true"},
            files={"file": ("real_import.xlsx", buf_real.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        )
        tr.record("Kịch bản 7 (BulkImport)", "Tạo tài khoản thật trả về 200", real_resp.status_code == 200)
        tr.record("Kịch bản 7 (BulkImport)", "Header Cache-Control chứa no-store", "no-store" in real_resp.headers.get("cache-control", ""))

        real_json = real_resp.json()
        real_summary = real_json.get("summary", {})
        tr.record("Kịch bản 7 (BulkImport)", "Tạo thành công 2 tài khoản mới", real_summary.get("created") == 2 and real_summary.get("valid") == 0)

        created_passwords = {}
        for row in real_json.get("rows", []):
            if row.get("status") == "created":
                created_passwords[row["email"]] = row.get("password")

        pwd1 = created_passwords.get(real_user1_email)
        tr.record("Kịch bản 7 (BulkImport)", "Mật khẩu tạm 12 ký tự không chứa ký tự nhầm lẫn", bool(pwd1 and len(pwd1) == 12 and not any(c in "0O1lI" for c in pwd1)))

        user1_client = httpx.Client(base_url=base_url, timeout=15.0)
        u1_login = auth_login(user1_client, real_user1_email, pwd1)
        tr.record("Kịch bản 7 (BulkImport)", "Đăng nhập được bằng mật khẩu tạm", u1_login.status_code == 200)

        u1_me = user1_client.get("/auth/me").json()
        tr.record(
            "Kịch bản 7 (BulkImport)",
            "Tài khoản mới có cờ must_change_password=True và vai trò user",
            u1_me.get("must_change_password") is True and u1_me.get("role") == "user"
        )

        u1_blocked = user1_client.get("/conversations")
        tr.record("Kịch bản 7 (BulkImport)", "Tài khoản mới bị chặn 403 trước khi đổi mật khẩu", u1_blocked.status_code == 403)

        # Lấy user id để dọn dẹp và kiểm tra daily_limit từ admin
        all_users = admin_client.get(f"{base_url}/admin/users").json()
        target_u1 = next((u for u in all_users if u["email"] == real_user1_email), None)
        tr.record(
            "Kịch bản 7 (BulkImport)",
            "Tài khoản mới có đúng daily_limit=30 trong hệ thống",
            bool(target_u1 and target_u1.get("daily_limit") == 30)
        )

        for u in all_users:
            if u["email"] in (real_user1_email, real_user2_email):
                users_to_cleanup.append(u["id"])

        # 7.6 Chạy lại cùng file: Toàn bộ dòng chuyển thành exists
        rerun_resp = admin_client.post(
            f"{base_url}/admin/users-import",
            data={"dry_run": "true"},
            files={"file": ("real_import.xlsx", buf_real.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        )
        rerun_summary = rerun_resp.json().get("summary", {})
        tr.record(
            "Kịch bản 7 (BulkImport)",
            "Chạy lại cùng file: toàn bộ dòng chuyển thành exists",
            rerun_summary.get("exists") == 2 and rerun_summary.get("created") == 0
        )

    finally:
        # ==========================================
        # DỌN SẠCH TẤT CẢ DỮ LIỆU THỬ NGHIỆM
        # ==========================================
        print("\n>> Dọn sạch dữ liệu người dùng thử nghiệm...")
        for uid in users_to_cleanup:
            try:
                purge_user_via_admin(admin_client, base_url, uid)
                print(f"  Đã xóa sạch user test ID: {uid}")
            except Exception as e:
                print(f"  Cảnh báo dọn dẹp user ID {uid}: {e}")

    return tr.print_summary()


def main():
    parser = argparse.ArgumentParser(description="Chạy Smoke Test kiểm tra toàn bộ hệ thống RAG Assistant.")
    parser.add_argument("--base-url", default=os.getenv("BASE_URL", "http://127.0.0.1:8000"), help="Địa chỉ server")
    parser.add_argument("--admin-email", default=os.getenv("ADMIN_EMAIL", DEFAULT_ADMIN_EMAIL), help="Email quản trị")
    parser.add_argument("--admin-password", default=os.getenv("ADMIN_PASSWORD", DEFAULT_ADMIN_PASSWORD), help="Mật khẩu quản trị")

    args = parser.parse_args()

    if not args.admin_email or not args.admin_password:
        print("LỖI: Chưa cung cấp thông tin tài khoản admin qua tham số (--admin-email, --admin-password) hoặc biến môi trường (ADMIN_EMAIL, ADMIN_PASSWORD).")
        sys.exit(1)

    success = run_smoke_tests(args.base_url.rstrip("/"), args.admin_email, args.admin_password)
    sys.exit(0 if success else 1)


if __name__ == "__main__":
    main()
